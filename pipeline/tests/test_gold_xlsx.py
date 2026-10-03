"""The formatted gold workbook (gold_labels.xlsx): layout, dropdowns, protection, text cells,
and the round trip through ``gold import``."""

from __future__ import annotations

import csv
import io
import json
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook
from test_gold import (
    MOVIE_LABELS,
    PREFILLED_MOVIE,
    STAMP,
    TV_LABELS,
    cand,
    sheet,
)
from wikimedia_fake import FakeWikimedia

from laminary_pipeline.annotation import validate_record
from laminary_pipeline.gold import columns as col
from laminary_pipeline.gold.__main__ import main as gold_main
from laminary_pipeline.gold.importer import (
    import_csv_text,
    import_sheet,
    import_xlsx,
)
from laminary_pipeline.gold.template import _count, lists_csv, template_csv, template_rows
from laminary_pipeline.gold.workbook import (
    FIRST_DATA_ROW,
    GROUPS,
    GUIDE_DOCX_NAME,
    HEADER_ROW,
    LABELS_SHEET,
    LISTS_SHEET,
    README_SHEET,
    XLSX_NAME,
    list_ranges,
    write_workbook,
)
from laminary_pipeline.ingest.http import HttpClient
from laminary_pipeline.ingest.paths import read_jsonl, write_jsonl_atomic
from laminary_pipeline.ingest.wikipedia import PlotFetcher


def _fixture_rows() -> list[dict[str, str]]:
    fetcher = PlotFetcher(HttpClient(FakeWikimedia(), None, min_interval={}),
                          clock=lambda: "2026-09-30T12:00:00Z")
    plots = {
        "Q9000001": fetcher.fetch({"qid": "Q9000001", "enwiki_title": "The Lantern Keeper",
                                   "media_type": "movie", "title": "The Lantern Keeper"}),
        "Q9000003": fetcher.fetch({"qid": "Q9000003", "enwiki_title": "Harbor Lights (TV series)",
                                   "media_type": "tv_series", "series_status": "ended"}),
    }
    selection = [
        {"qid": "Q9000001", "title": "The Lantern Keeper", "year": 1994, "media_type": "movie",
         "tmdb_id": 90001},
        {"qid": "Q9000003", "title": "Harbor Lights", "year": 2013, "media_type": "tv_series",
         "tmdb_id": 90003, "double_label": True},
    ]
    rows, skipped = template_rows(selection, plots)
    assert skipped == [] and len(rows) == 3
    return rows


def _labels(ws: Any) -> list[list[Any]]:
    return [list(r) for r in ws.iter_rows(values_only=True)]


@pytest.fixture
def book(tmp_path: Path) -> tuple[Path, list[dict[str, str]]]:
    rows = _fixture_rows()
    return write_workbook(tmp_path / XLSX_NAME, rows), rows


def test_labels_sheet_has_the_csv_rows_and_columns(book) -> None:
    path, rows = book
    wb = load_workbook(path)
    assert wb.sheetnames == [README_SHEET, LABELS_SHEET, LISTS_SHEET]
    ws = wb[LABELS_SHEET]
    values = _labels(ws)
    assert values[HEADER_ROW - 1] == col.ALL_COLUMNS
    # the CSV without its '#' help row, cell for cell (blank labeler cells are empty)
    csv_rows = list(csv.reader(io.StringIO(template_csv(rows))))[2:]
    got = [[v if v is not None else "" for v in r] for r in values[FIRST_DATA_ROW - 1:]]
    assert got == csv_rows
    assert not any(str(v or "").startswith("#") for r in values for v in r)
    # group row names each group once, over its first column
    group_row = values[0]
    for name, cols in GROUPS:
        assert group_row[col.ALL_COLUMNS.index(cols[0])] == name
    assert [c for g in GROUPS for c in g[1]] == col.ALL_COLUMNS


def test_prefilled_cells_are_text_locked_and_shaded(book) -> None:
    path, rows = book
    ws = load_workbook(path)[LABELS_SHEET]
    idx = {c: i + 1 for i, c in enumerate(col.ALL_COLUMNS)}
    r = FIRST_DATA_ROW
    for c in ("year", "source_retrieved_at", "tmdb_id", "word_count", "label_slot"):
        cell = ws.cell(row=r, column=idx[c])
        assert cell.data_type == "s" and cell.number_format == "@", c
        assert isinstance(cell.value, str)
    assert ws.cell(row=r, column=idx["source_retrieved_at"]).value == "2026-09-30T12:00:00Z"
    assert ws.protection.sheet and not ws.protection.password
    for c in col.ALL_COLUMNS:
        cell = ws.cell(row=r, column=idx[c])
        assert cell.protection.locked == (c not in col.LABELER_COLS), c
        white = cell.fill.fgColor.rgb.endswith("FFFFFF")
        assert white == (c in col.LABELER_COLS), c
    assert ws.freeze_panes == "F3"  # header rows and qid..label_slot stay in view
    # links: the text file and a single revision link are clickable
    assert ws.cell(row=r, column=idx["summary_text_file"]).hyperlink.target == "texts/Q9000001.txt"
    link = ws.cell(row=r, column=idx["wikipedia_revision_link"])
    assert link.hyperlink.target == rows[0]["wikipedia_revision_link"]
    # header help lives in notes on the header cells
    note = ws.cell(row=HEADER_ROW, column=idx["summary_text_file"]).comment
    assert "READ THIS FILE ONLY" in note.text


def test_slot_two_rows_are_marked(book) -> None:
    path, rows = book
    ws = load_workbook(path)[LABELS_SHEET]
    slot = col.ALL_COLUMNS.index("label_slot") + 1
    fills = {}
    for i, row in enumerate(rows):
        cell = ws.cell(row=FIRST_DATA_ROW + i, column=slot)
        fills[row["label_slot"]] = (cell.fill.fgColor.rgb, cell.font.b)
    assert fills["1"] != fills["2"] and fills["2"][1] is True


def test_every_fixed_value_column_has_a_range_dropdown(book) -> None:
    path, rows = book
    ws = load_workbook(path)[LABELS_SHEET]
    idx = {c: i + 1 for i, c in enumerate(col.ALL_COLUMNS)}
    from openpyxl.utils import get_column_letter

    by_col: dict[str, Any] = {}
    for dv in ws.data_validations.dataValidation:
        for rng in dv.sqref.ranges:
            by_col[get_column_letter(rng.min_col)] = (dv, rng)
    ranges = list_ranges()
    last = FIRST_DATA_ROW + len(rows) - 1
    for c in col.LABELER_COLS:
        dd = col.dropdown_for_column(c)
        entry = by_col.get(get_column_letter(idx[c]))
        if dd is None and c not in col.ARC_POINT_COLS:
            continue
        assert entry is not None, c
        dv, rng = entry
        assert (rng.min_row, rng.max_row) == (FIRST_DATA_ROW, last)
        assert dv.showErrorMessage and dv.errorStyle == "stop", c
        if dd is not None:
            assert dv.type == "list" and dv.formula1 == ranges[dd], c
            assert dv.formula1.startswith(f"{LISTS_SHEET}!$")  # a range, not an inline list
        else:
            assert dv.type == "decimal" and (dv.formula1, dv.formula2) == ("-1", "1")
    # the ranges cover exactly the Lists tab's values, which equal the CSV's
    wb = load_workbook(path)
    lists = _labels(wb[LISTS_SHEET])
    assert [[v or "" for v in r] for r in lists] == list(csv.reader(io.StringIO(lists_csv())))
    assert wb[LISTS_SHEET].protection.sheet
    for name, values in col.dropdown_lists().items():
        assert ranges[name].endswith(f"${len(values) + 1}")


def test_readme_names_the_guide_and_the_sheets_caveat(book) -> None:
    path, _ = book
    text = "\n".join(str(v) for r in _labels(load_workbook(path)[README_SHEET]) for v in r if v)
    assert GUIDE_DOCX_NAME in text
    assert "Data > Protect sheets and ranges" in text
    assert 5 <= sum(1 for line in text.splitlines() if line[:2] in {f"{i}." for i in range(9)})


def test_formula_like_values_stay_text(tmp_path: Path) -> None:
    rows = _fixture_rows()
    rows[0] = {**rows[0], "title": "=HYPERLINK(\"x\")", "year": "1994"}
    ws = load_workbook(write_workbook(tmp_path / "f.xlsx", rows))[LABELS_SHEET]
    cell = ws.cell(row=FIRST_DATA_ROW, column=col.ALL_COLUMNS.index("title") + 1)
    assert cell.data_type == "s" and cell.value == "=HYPERLINK(\"x\")"


ARC_POINTS = ["0", "-0.2", "-0.4", "-0.6", "-0.7", "-0.6", "-0.3", "0", "0.3", "0.5", "0.6"]


def _fill_all(row: dict[str, str]) -> dict[str, str]:
    """Every labeler cell filled with a valid value (skip_reason left blank: it replaces the
    labels)."""
    slot = row["label_slot"]
    labels = dict(MOVIE_LABELS, labeler_id="L01" if slot == "1" else "L02")
    labels.update(zip(col.ARC_POINT_COLS, ARC_POINTS, strict=True))
    return labels


def _fill_xlsx(path: Path, rows: list[dict[str, str]], numbers: bool) -> None:
    wb = load_workbook(path)
    ws = wb[LABELS_SHEET]
    idx = {c: i + 1 for i, c in enumerate(col.ALL_COLUMNS)}
    for i, row in enumerate(rows):
        for c, v in _fill_all(row).items():
            # a spreadsheet stores typed arc points as numbers
            ws.cell(row=FIRST_DATA_ROW + i, column=idx[c]).value = (
                float(v) if numbers and c in col.ARC_POINT_COLS else v)
    wb.save(path)


def test_filled_workbook_round_trips_like_the_csv(book, tmp_path: Path) -> None:
    path, rows = book
    _fill_xlsx(path, rows, numbers=True)
    from_xlsx = import_xlsx(path, annotated_at=STAMP)
    assert from_xlsx.errors == [], "\n".join(map(str, from_xlsx.errors))
    assert len(from_xlsx.records) == len(rows) == 3
    assert [validate_record(r) for r in from_xlsx.records] == [[], [], []]
    filled = [{**r, **_fill_all(r)} for r in rows]
    from_csv = import_csv_text(sheet(*filled), annotated_at=STAMP)
    assert from_csv.errors == []
    assert from_xlsx.records == from_csv.records
    rec = from_xlsx.records[0]
    assert rec["title"]["release_year"] == 1994 and rec["provenance"]["sources"][0][
        "retrieved_at"] == "2026-09-30T12:00:00Z"
    assert rec["layers"]["structural_skeleton"]["arc_points"] == [float(p) for p in ARC_POINTS]
    assert import_sheet(path, annotated_at=STAMP).records == from_xlsx.records


def test_workbook_errors_name_the_spreadsheet_row(book) -> None:
    path, rows = book
    _fill_xlsx(path, rows, numbers=False)
    wb = load_workbook(path)
    ws = wb[LABELS_SHEET]
    ws.cell(row=FIRST_DATA_ROW, column=col.ALL_COLUMNS.index("plot_rebirth") + 1).value = "N"
    wb.save(path)
    result = import_xlsx(path, annotated_at=STAMP)
    assert str(result.errors[0]).startswith(f"Row {FIRST_DATA_ROW} (The Lantern Keeper)")


def test_csv_exported_from_the_workbook_keeps_its_group_row() -> None:
    """A CSV download of the Sheets copy starts with the group row; the header is found below
    it and row numbers stay the spreadsheet's."""
    text = sheet({**PREFILLED_MOVIE, **MOVIE_LABELS}, help_row=False)
    groups = ",".join(["Title"] + [""] * (len(col.ALL_COLUMNS) - 1))
    result = import_csv_text(groups + "\n" + text, annotated_at=STAMP)
    assert result.errors == [] and len(result.records) == 1
    bad = sheet({**PREFILLED_MOVIE, **MOVIE_LABELS, "primary_plot": ""}, help_row=False)
    result = import_csv_text(groups + "\n" + bad, annotated_at=STAMP)
    assert str(result.errors[0]).startswith("Row 3 ")


def test_cli_template_writes_the_workbook_and_imports_it(tmp_path: Path) -> None:
    data = tmp_path / "data"
    fetcher = PlotFetcher(HttpClient(FakeWikimedia(), None, min_interval={}),
                          clock=lambda: "2026-09-30T12:00:00Z")
    c = {**cand(9000001, "movie", "english", "The Lantern Keeper"), "qid": "Q9000001",
         "enwiki_title": "The Lantern Keeper", "bucket": "film:english", "bucket_rank": 1}
    write_jsonl_atomic(data / "pilot_candidates.jsonl", [c])
    (data / "plots").mkdir(parents=True)
    (data / "plots" / "Q9000001.json").write_text(json.dumps(fetcher.fetch(c)))
    log: list[str] = []
    assert gold_main(["--data-dir", str(data), "select", "--n", "5"], log=log.append) == 0
    # the selection (with the guesses) stays out of the upload folder
    assert (data / "gold_internal" / "gold_selection.jsonl").exists()
    assert not (data / "gold" / "gold_selection.jsonl").exists()
    assert gold_main(["--data-dir", str(data), "template"], log=log.append) == 0
    book_path = data / "gold" / XLSX_NAME
    assert book_path.exists() and (data / "gold" / "gold_labels_template.csv").exists()
    assert sorted(p.name for p in (data / "gold").iterdir()) == [
        "gold_labels.xlsx", "gold_labels_lists.csv", "gold_labels_readme.csv",
        "gold_labels_template.csv", "texts"]
    rows = [r for r in csv.DictReader(io.StringIO(
        (data / "gold" / "gold_labels_template.csv").read_text()))][1:]
    _fill_xlsx(book_path, rows, numbers=True)
    code = gold_main(["--data-dir", str(data), "import", str(book_path), "--labeled-at", STAMP],
                     log=log.append)
    assert code == 0, "\n".join(log)
    records = list(read_jsonl(data / "gold" / "gold_labels.jsonl"))
    assert len(records) == len(rows) == 2  # one title, both slots (fewer titles than 25)
    assert [validate_record(r) for r in records] == [[], []]


def test_template_points_at_a_selection_left_in_the_old_place(tmp_path: Path) -> None:
    data = tmp_path / "data"
    write_jsonl_atomic(data / "gold" / "gold_selection.jsonl", [{"qid": "Q1"}])
    log: list[str] = []
    assert gold_main(["--data-dir", str(data), "template"], log=log.append) == 2
    assert "move it to" in log[-1] and "gold_internal" in log[-1]


def test_unit_counts_are_singular_for_one() -> None:
    assert _count(1, "season") == "1 season" and _count(3, "season") == "3 seasons"
    assert _count(1, "year") == "1 year" and _count(2, "year") == "2 years"


def test_tv_labels_still_import_from_a_workbook(book) -> None:
    """Arc points with no arc_shape, lower-case y/n and schema keys work in the workbook too."""
    path, rows = book
    wb = load_workbook(path)
    ws = wb[LABELS_SHEET]
    idx = {c: i + 1 for i, c in enumerate(col.ALL_COLUMNS)}
    for c, v in TV_LABELS.items():
        ws.cell(row=FIRST_DATA_ROW + 1, column=idx[c]).value = v
    wb.save(path)
    result = import_xlsx(path, annotated_at=STAMP)
    assert result.errors == [] and len(result.records) == 1
    assert result.records[0]["layers"]["structural_skeleton"]["emotional_arc"]["method"] == (
        "derived")
